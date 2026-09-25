import os
import sys
import json
import time
import glob
import ctypes
import threading
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

try:
    import winsound
except ImportError:
    winsound = None

# ==========================================
# 1. FRONT 전용 모델 설정
# ==========================================
MODEL_CONFIG = {
    'S-FRONT': 'MPL02916AD',
    'R-FRONT': 'MPL02926AD'
}

CODE_TO_MODEL = {v: k for k, v in MODEL_CONFIG.items()}
DEFAULT_PASSWORD = "123456"
MAX_ITEMS_PER_BOX = 10
MAX_BOXES_PER_PALLET = 12

def get_base_dir():
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))

BASE_DIR = get_base_dir()
COUNT_FILE = os.path.join(BASE_DIR, "counts_front.json")
STATE_FILE = os.path.join(BASE_DIR, "pallet_state_front.json")

FILE_ATTRIBUTE_NORMAL = 0x80
FILE_ATTRIBUTE_HIDDEN = 0x02

def unhide_file(filepath):
    try:
        if os.name == 'nt' and os.path.exists(filepath):
            ctypes.windll.kernel32.SetFileAttributesW(str(filepath), FILE_ATTRIBUTE_NORMAL)
    except Exception:
        pass

def hide_file(filepath):
    try:
        if os.name == 'nt' and os.path.exists(filepath):
            ctypes.windll.kernel32.SetFileAttributesW(str(filepath), FILE_ATTRIBUTE_HIDDEN)
    except Exception:
        pass

def get_quarter_filename(model_name, dt=None):
    if dt is None:
        dt = datetime.now()
    year_2d = dt.strftime("%y")
    quarter = (dt.month - 1) // 3 + 1
    safe_model = model_name.replace('-', '_')
    return f"Y{year_2d}_{quarter}Q_{safe_model}.xlsx"

LANG_PACK = {
    "한국어": {
        "title": "QR SCAN STATION [FRONT]",
        "pw_setting": "⚙ 비밀번호 설정",
        "tab_scan": "  QR Scan  ",
        "tab_grouping": "  Grouping  ",
        "tab_recode": "  Re-code  ",
        "model_label": "모델",
        "code_label": "인식코드",
        "last_scan": "마지막 스캔",
        "input_guide": "바코드 스캔 입력 (어느 화면에서나 스캔 가능)",
        "reset_btn": "RESET (카운터 초기화)",
        "manager_btn": "MANAGER MODE",
        "manager_btn_on": "MANAGER MODE [ON]",
        "pending_status": "미그룹 스캔 {count}건 – Label QR 대기 중",
        "pallet_status": "현재 팔레트: {pallet} ({boxes}/{max_b} 박스)",
        "record_header": "{model} 기록",
        "grouping_header": "{model} Pallet - Label Grouping 현황",
        "th_pallet": "Pallet Label QR",
        "th_box_seq": "박스 번호",
        "th_day": "DAY",
        "th_time": "TIME",
        "th_label": "Label QR",
        "th_dmc": "DMC",
        "th_judgment": "JUDGMENT",
        "th_content": "Content",
        "filter_day": "조회 일자:",
        "filter_time": "시간대:",
        "search_btn": "🔍 검색",
        "save_btn": "💾 Save (엑셀 저장)",
        "box_complete": "[Box Grouping Done: {count} pcs]",
        "dup_scan_tag": "[중복 스캔]",
        "sorting_title": "⚠️ Sorting 필요 제품 경고",
        "sorting_msg": "[알림: Sorting 필요 제품]\n\nDMC Code: {code}\n\n해당 제품은 Sorting 대상 리스트에 등록되어 있습니다.\n바코드를 별도로 격리한 뒤 [Enter] 키를 누르세요.",
        "ng_model_title": "⚠️ NG - 모델 불일치",
        "ng_model_msg": "[NG: 선택 모델과 바코드 코드가 일치하지 않습니다]\n\n현재 선택 모델: {model} ({target})\n스캔된 코드: {code}\n\n관리자 비밀번호 6자리를 입력하여 해제하세요.",
        "ng_pallet_model_title": "⚠️ NG - Pallet QR 모델 불일치",
        "ng_pallet_model_msg": "[NG: Pallet QR 모델 코드가 일치하지 않습니다]\n\n현재 선택 모델: {model} ({target})\n스캔 Pallet QR: {code}\n\n올바른 Pallet QR을 준비한 뒤 관리자 비밀번호로 해제하세요.",
        "ng_label_dup_title": "⚠️ Label QR NG - 중복 스캔",
        "ng_label_dup_msg": "[Label QR NG: 이미 사용된 Label QR입니다]\n\n스캔 Label QR: {code}...\n이미 등록/포장 완료된 중복 라벨입니다.\n\n관리자 비밀번호 6자리를 입력하여 해제하세요.",
        "ng_group_title": "⚠️ Grouping NG - 수량 불일치",
        "ng_group_msg": "[Grouping NG: 단품 수량과 Label 포장 수량 불일치]\n\nLabel QR 지정 수량: {expected}개\n현재 스캔된 단품 수량: {current}개\n\n수량이 일치하지 않아 묶음을 진행할 수 없습니다.\n관리자 비밀번호 6자리를 입력하여 해제하세요.",
        "ng_limit_title": "⚠️ NG - Label QR 누락",
        "ng_limit_msg": "[NG 발생: Label QR 누락]\n\n단품이 이미 {max_cnt}개 모두 스캔되었습니다.\n11번째 단품은 기록되지 않습니다.\nLabel QR을 먼저 스캔하여 박스 묶음을 완료하십시오.\n\n관리자 비밀번호 6자리를 입력하여 해제하세요.",
        "ng_mgr_err_title": "⚠️ NG - 관리자 모드 오류",
        "ng_mgr_err_msg": "[NG: 관리자 모드가 아닌 일반 모드에서 QR 리딩 필요]\n\n스캔 바코드: {code}\n신규 제품은 일반 모드에서 등록해야 합니다.\n해당 스캔은 기록되지 않습니다.\n\n관리자 비밀번호 6자리를 입력하여 해제하세요.",
        "ng_dup_title": "🚫 QR NG - 중복 바코드 감지",
        "ng_dup_msg": "[QR NG 발생: 이미 스캔된 바코드입니다]\n\n스캔 바코드: {code}\n해당 제품 및 연결된 박스 헤더가 NG로 변경되었습니다.\n\n관리자 비밀번호 6자리를 입력하여 해제하세요.",
        "ng_pallet_mid_title": "⚠️ NG - Pallet 리딩 시점 오류",
        "ng_pallet_mid_msg": "[NG: 단품 스캔 도중에는 Pallet QR을 리딩할 수 없습니다]\n\n현재 {count}개의 단품이 스캔 중입니다.\n10개 단품 및 Label QR 스캔을 완료한 후 Pallet QR을 리딩하세요.",
        "ng_pallet_limit_title": "🚫 NG - Pallet QR 누락 (13박스 초과)",
        "ng_pallet_limit_msg": "[NG: Pallet QR 리딩 누락]\n\n이미 12개 박스가 채워졌습니다.\n새 Pallet QR을 리딩하지 않고 13번째 이상 박스를 진행할 수 없습니다.\n\n관리자 비밀번호를 입력하여 해제하세요.",
        "pallet_popup_title": "Pallet QR 스캔 대기",
        "pallet_popup_msg": "12개 박스 포장이 완료되었습니다.\n새로운 Pallet QR을 스캔해주세요.",
        "unlock_btn": "확인 및 잠금 해제",
        "confirm_btn": "확인 (Enter)",
        "pw_err": "비밀번호가 올바르지 않습니다."
    },
    "English": {
        "title": "QR SCAN STATION [FRONT]",
        "pw_setting": "⚙ Password Setting",
        "tab_scan": "  QR Scan  ",
        "tab_grouping": "  Grouping  ",
        "tab_recode": "  Re-code  ",
        "model_label": "Model",
        "code_label": "Part No",
        "last_scan": "Last Scan",
        "input_guide": "Barcode Scan Input (Focus anywhere)",
        "reset_btn": "RESET (Clear Counter)",
        "manager_btn": "MANAGER MODE",
        "manager_btn_on": "MANAGER MODE [ON]",
        "pending_status": "Ungrouped: {count} pcs – Waiting for Label QR",
        "pallet_status": "Current Pallet: {pallet} ({boxes}/{max_b} Boxes)",
        "record_header": "{model} Records",
        "grouping_header": "{model} Pallet - Box Grouping Overview",
        "th_pallet": "Pallet Label QR",
        "th_box_seq": "Box Seq",
        "th_day": "DAY",
        "th_time": "TIME",
        "th_label": "Label QR",
        "th_dmc": "DMC",
        "th_judgment": "JUDGMENT",
        "th_content": "Content",
        "filter_day": "Date Range:",
        "filter_time": "Time Range:",
        "search_btn": "🔍 Search",
        "save_btn": "💾 Save (Excel Export)",
        "box_complete": "[Box Grouping Done: {count} pcs]",
        "dup_scan_tag": "[Duplicate Scan]",
        "sorting_title": "⚠️ Sorting Required Alert",
        "sorting_msg": "[Alert: Sorting Required Product]\n\nDMC Code: {code}\n\nThis product is registered in the Sorting list.\nIsolate the part and press [Enter] to continue.",
        "ng_model_title": "⚠️ NG - Model Mismatch",
        "ng_model_msg": "[NG: Scanned barcode does not match selected model]\n\nSelected Model: {model} ({target})\nScanned Code: {code}\n\nEnter 6-digit Admin Password to unlock.",
        "ng_pallet_model_title": "⚠️ NG - Pallet Model Mismatch",
        "ng_pallet_model_msg": "[NG: Pallet QR model code does not match]\n\nSelected Model: {model} ({target})\nScanned Pallet QR: {code}\n\nEnter 6-digit Admin Password to unlock.",
        "ng_label_dup_title": "⚠️ Label QR NG - Duplicate Label",
        "ng_label_dup_msg": "[Label QR NG: This Label QR is already used]\n\nScanned Label: {code}...\nDuplicate box label detected.\n\nEnter 6-digit Admin Password to unlock.",
        "ng_group_title": "⚠️ Grouping NG - Quantity Mismatch",
        "ng_group_msg": "[Grouping NG: Scanned quantity does not match Label quantity]\n\nLabel Target Qty: {expected} pcs\nCurrently Scanned Qty: {current} pcs\n\nCannot proceed with grouping.\nEnter 6-digit Admin Password to unlock.",
        "ng_limit_title": "⚠️ NG - Label QR Missing",
        "ng_limit_msg": "[NG: Label QR Missing]\n\nAlready reached maximum capacity ({max_cnt} pcs).\n11th item is rejected and not saved.\nScan Label QR first to complete the box.\n\nEnter 6-digit Admin Password to unlock.",
        "ng_mgr_err_title": "⚠️ NG - Manager Mode Error",
        "ng_mgr_err_msg": "[NG: New QR must be scanned in Normal Mode]\n\nScanned Barcode: {code}\nNew parts cannot be added under Manager Mode.\nScan discarded.\n\nEnter 6-digit Admin Password to unlock.",
        "ng_dup_title": "🚫 QR NG - Duplicate Part Detected",
        "ng_dup_msg": "[QR NG: Duplicate part barcode detected]\n\nScanned Barcode: {code}\nThis part and associated Box Header are marked as NG.\n\nEnter 6-digit Admin Password to unlock.",
        "ng_pallet_mid_title": "⚠️ NG - Invalid Pallet Scan Timing",
        "ng_pallet_mid_msg": "[NG: Cannot scan Pallet QR while box packing is in progress]\n\nCurrently {count} items are pending.\nFinish 10 items and Label QR before scanning Pallet QR.",
        "ng_pallet_limit_title": "🚫 NG - Pallet QR Missing (Exceeded 12 Boxes)",
        "ng_pallet_limit_msg": "[NG: Pallet QR Missing]\n\n12 boxes are already filled.\nCannot pack 13th box without scanning a new Pallet QR.\n\nEnter Admin Password to unlock.",
        "pallet_popup_title": "Waiting for Pallet QR",
        "pallet_popup_msg": "12 boxes completed on current pallet.\nPlease scan new Pallet QR.",
        "unlock_btn": "Confirm & Unlock",
        "confirm_btn": "Confirm (Enter)",
        "pw_err": "Incorrect Password."
    },
    "Polski": {
        "title": "QR SCAN STATION [FRONT]",
        "pw_setting": "⚙ Ustawienie hasła",
        "tab_scan": "  Skan QR  ",
        "tab_grouping": "  Grupowanie  ",
        "tab_recode": "  Re-code  ",
        "model_label": "Model",
        "code_label": "Kod części",
        "last_scan": "Ostatni skan",
        "input_guide": "Wejście skanera (skanuj w dowolnym miejscu)",
        "reset_btn": "RESET (Zeruj licznik)",
        "manager_btn": "TRYB MENEDŻERA",
        "manager_btn_on": "TRYB MENEDŻERA [ON]",
        "pending_status": "Oczekujące: {count} szt. – Oczekiwanie na Label QR",
        "pallet_status": "Bieżąca paleta: {pallet} ({boxes}/{max_b} pudełek)",
        "record_header": "{model} Historia",
        "grouping_header": "{model} Przegląd grupowania Paleta - Pudełko",
        "th_pallet": "Pallet Label QR",
        "th_box_seq": "Nr pudełka",
        "th_day": "DZIEŃ",
        "th_time": "CZAS",
        "th_label": "Label QR",
        "th_dmc": "DMC",
        "th_judgment": "STATUS",
        "th_content": "Treść",
        "filter_day": "Zakres dat:",
        "filter_time": "Przedział czasu:",
        "search_btn": "🔍 Szukaj",
        "save_btn": "💾 Zapisz (Eksport Excel)",
        "box_complete": "[Pakiet ukończony: {count} szt.]",
        "dup_scan_tag": "[Duplikat skanu]",
        "sorting_title": "⚠️ Wymagane sortowanie",
        "sorting_msg": "[Uwaga: Wymagane sortowanie produktu]\n\nKod DMC: {code}\n\nTen produkt znajduje się na liście sortowania.\nOdizoluj część i naciśnij [Enter], aby kontynuować.",
        "ng_model_title": "⚠️ NG - Niezgodność modelu",
        "ng_model_msg": "[NG: Zeskanowany kod nie pasuje do wybranego modelu]\n\nWybrany model: {model} ({target})\nKod: {code}\n\nWprowadź 6-cyfrowe hasło administratora, aby odblokować.",
        "ng_pallet_model_title": "⚠️ NG - Niezgodność modelu palety",
        "ng_pallet_model_msg": "[NG: Kod modelu na etykiecie palety nie pasuje]\n\nWybrany model: {model} ({target})\nPaleta: {code}\n\nWprowadź 6-cyfrowe hasło administratora.",
        "ng_label_dup_title": "⚠️ Label QR NG - Duplikat etykiety",
        "ng_label_dup_msg": "[Label QR NG: Ta etykieta została już użyta]\n\nZeskanowana etykieta: {code}...\nWykryto duplikat etykiety pudełka.\n\nWprowadź 6-cyfrowe hasło administratora, aby odblokować.",
        "ng_group_title": "⚠️ Grouping NG - Niezgodność ilości",
        "ng_group_msg": "[Grouping NG: Ilość sztuk nie zgadza się z etykietą]\n\nIlość na etykiecie: {expected} szt.\nZeskanowano: {current} szt.\n\nNie można utworzyć grupy.\nWprowadź 6-cyfrowe hasło administratora, aby odblokować.",
        "ng_limit_title": "⚠️ NG - Brak Label QR",
        "ng_limit_msg": "[NG: Brak Label QR]\n\nOsiągnięto limit pudełka ({max_cnt} szt.).\n11. element nie został zapisany.\nZeskanuj najpierw Label QR.\n\nWprowadź 6-cyfrowe hasło administratora, aby odblokować.",
        "ng_mgr_err_title": "⚠️ NG - Błąd trybu menedżera",
        "ng_mgr_err_msg": "[NG: Nowe części należy skanować w trybie standardowym]\n\nZeskanowany kod: {code}\nNowy element został odrzucony.\n\nWprowadź 6-cyfrowe hasło administratora, aby odblokować.",
        "ng_dup_title": "🚫 QR NG - Wykryto zduplikowany element",
        "ng_dup_msg": "[QR NG: Kod tego elementu został 이미 이전 기록에 있습니다]\n\nZeskanowany kod: {code}\nTen element i nagłówek partii oznaczono jako NG.\n\nWprowadź 6-cyfrowe hasło administratora, aby odblokować.",
        "ng_pallet_mid_title": "⚠️ NG - Błędny moment skanowania palety",
        "ng_pallet_mid_msg": "[NG: Nie można skanować kodu palety podczas pakowania pudełka]\n\nObecnie oczekuje {count} elementów.\nZakończ 10 sztuk i Label QR przed zeskanowaniem palety.",
        "ng_pallet_limit_title": "🚫 NG - Brak kodu palety (Przekroczono 12 pudełek)",
        "ng_pallet_limit_msg": "[NG: Wymagany nowy kod palety]\n\nZapakowano już 12 pudełek.\nNie można kontynuować 13. pudełka bez nowej palety.\n\nWprowadź hasło administratora.",
        "pallet_popup_title": "Oczekiwanie na kod palety",
        "pallet_popup_msg": "Ukończono 12 pudełek na palecie.\nZeskanuj kod nowej palety.",
        "unlock_btn": "Potwierdź i odblokuj",
        "confirm_btn": "Potwierdź (Enter)",
        "pw_err": "Nieprawidłowe hasło."
    }
}

BG_MAIN = "#1a1f26"
BG_PANEL = "#222731"
BG_INPUT = "#15181e"
TEXT_COLOR = "#e1e4ea"
TEXT_MUTED = "#8b949e"
ACCENT_YELLOW = "#f59f00"
COLOR_OK = "#28a745"
COLOR_NG = "#dc3545"


class QRScanStationApp:
    def __init__(self, root):
        self.root = root
        self.root.title("QR SCAN STATION [FRONT]")
        self.root.geometry("1420x820")
        self.root.minsize(1240, 740)
        self.root.configure(bg=BG_MAIN)

        self.current_lang = tk.StringVar(value="한국어")
        self.current_model = tk.StringVar(value='S-FRONT')
        self.admin_password = DEFAULT_PASSWORD
        self.model_session_id = 0

        self.is_manager_mode = False
        self.active_popup = None
        self.pallet_wait_popup = None

        self.last_scanned_code = ""
        self.last_scanned_time = 0.0
        self.auto_submit_timer = None

        self.model_counts = self.load_model_counts()
        self.pallet_state = self.load_pallet_state()

        self.scanned_history_by_model = {m: set() for m in MODEL_CONFIG}
        self.scanned_label_by_model = {m: set() for m in MODEL_CONFIG}
        self.sorting_list_by_model = {m: set() for m in MODEL_CONFIG}

        self.pending_items = []
        self.pending_tree_ids = []
        self.file_lock = threading.Lock()
        self.global_scan_buffer = []

        self.setup_custom_styles()
        self.setup_ui()
        self.setup_global_key_listener()
        self.on_model_changed()

    def t(self, key, **kwargs):
        pack = LANG_PACK.get(self.current_lang.get(), LANG_PACK["한국어"])
        text = pack.get(key, "")
        if kwargs:
            return text.format(**kwargs)
        return text

    def play_alarm_sound(self):
        def _beep():
            if winsound:
                for _ in range(3):
                    winsound.Beep(1000, 350)
                    time.sleep(0.08)
        threading.Thread(target=_beep, daemon=True).start()

    def load_model_counts(self):
        default_counts = {m: {"total": 0, "ok": 0, "ng": 0} for m in MODEL_CONFIG}
        if os.path.exists(COUNT_FILE):
            try:
                with open(COUNT_FILE, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    for m in default_counts:
                        if m in data:
                            default_counts[m] = data[m]
                    return default_counts
            except Exception:
                return default_counts
        return default_counts

    def save_model_counts(self):
        try:
            with open(COUNT_FILE, "w", encoding="utf-8") as f:
                json.dump(self.model_counts, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

    def load_pallet_state(self):
        default_state = {m: {"current_pallet": "", "box_count": 0} for m in MODEL_CONFIG}
        if os.path.exists(STATE_FILE):
            try:
                with open(STATE_FILE, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    for m in default_state:
                        if m in data:
                            default_state[m] = data[m]
                    return default_state
            except Exception:
                return default_state
        return default_state

    def save_pallet_state(self):
        try:
            with open(STATE_FILE, "w", encoding="utf-8") as f:
                json.dump(self.pallet_state, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

    def setup_custom_styles(self):
        style = ttk.Style()
        style.theme_use("clam")

        style.configure("Dark.TNotebook", background=BG_MAIN, borderwidth=0)
        style.configure("Dark.TNotebook.Tab", background="#2a2f3a", foreground=TEXT_MUTED,
                        font=("맑은 고딕", 10, "bold"), padding=[15, 5])
        style.map("Dark.TNotebook.Tab",
                  background=[("selected", BG_PANEL)],
                  foreground=[("selected", "#ffffff")])

        style.configure("Dark.TCombobox", 
                        fieldbackground=BG_INPUT, 
                        background="#303642", 
                        foreground="#ffffff", 
                        arrowcolor="#ffffff",
                        darkcolor=BG_INPUT, 
                        lightcolor=BG_INPUT)

        style.configure("Dark.Treeview",
                        background="#1e232d",
                        foreground=TEXT_COLOR,
                        fieldbackground="#1e232d",
                        rowheight=26,
                        font=("맑은 고딕", 9))
        style.configure("Dark.Treeview.Heading",
                        background="#282e3a",
                        foreground="#adb5bd",
                        font=("맑은 고딕", 9, "bold"),
                        relief="flat")
        style.map("Dark.Treeview",
                  background=[("selected", "#2b5278")],
                  foreground=[("selected", "#ffffff")])
        style.map("Dark.Treeview.Heading",
                  background=[("active", "#343c4c")])

    def setup_ui(self):
        header_frame = tk.Frame(self.root, bg=BG_MAIN, height=45)
        header_frame.pack(fill=tk.X, padx=20, pady=(10, 4))

        tk.Label(header_frame, text="QR  SCAN  STATION  [FRONT]", font=("Arial", 12, "bold"), 
                 fg=TEXT_COLOR, bg=BG_MAIN).pack(side=tk.LEFT, padx=(0, 15))

        self.model_combo = ttk.Combobox(
            header_frame, 
            textvariable=self.current_model, 
            values=list(MODEL_CONFIG.keys()), 
            state="readonly", 
            font=("맑은 고딕", 10, "bold"), 
            width=13,
            style="Dark.TCombobox"
        )
        self.model_combo.pack(side=tk.LEFT)
        self.model_combo.bind("<<ComboboxSelected>>", self.on_model_changed)

        self.lang_combo = ttk.Combobox(
            header_frame,
            textvariable=self.current_lang,
            values=["한국어", "English", "Polski"],
            state="readonly",
            font=("맑은 고딕", 9, "bold"),
            width=9,
            style="Dark.TCombobox"
        )
        self.lang_combo.pack(side=tk.RIGHT, padx=(10, 0))
        self.lang_combo.bind("<<ComboboxSelected>>", self.on_language_changed)

        self.lbl_lang_icon = tk.Label(header_frame, text="🌐", font=("맑은 고딕", 12), bg=BG_MAIN, fg=TEXT_COLOR)
        self.lbl_lang_icon.pack(side=tk.RIGHT)

        self.btn_pw = tk.Button(
            header_frame, text=self.t("pw_setting"), command=self.change_password_dialog,
            bg="#2c323d", fg=TEXT_COLOR, activebackground="#3a4250", activeforeground="#ffffff",
            relief="flat", font=("맑은 고딕", 9), padx=10, pady=3, cursor="hand2"
        )
        self.btn_pw.pack(side=tk.RIGHT, padx=(0, 15))

        self.notebook = ttk.Notebook(self.root, style="Dark.TNotebook")
        self.notebook.pack(fill=tk.BOTH, expand=True, padx=20, pady=(0, 10))

        # 3개 탭 구성 (QR Scan, Grouping, Re-code)
        self.tab_scan = tk.Frame(self.notebook, bg=BG_MAIN)
        self.notebook.add(self.tab_scan, text=self.t("tab_scan"))

        self.tab_grouping = tk.Frame(self.notebook, bg=BG_MAIN)
        self.notebook.add(self.tab_grouping, text=self.t("tab_grouping"))

        self.tab_recode = tk.Frame(self.notebook, bg=BG_MAIN)
        self.notebook.add(self.tab_recode, text=self.t("tab_recode"))

        self.build_scan_tab()
        self.build_grouping_tab()
        self.build_recode_tab()

    def build_scan_tab(self):
        main_frame = tk.Frame(self.tab_scan, bg=BG_MAIN)
        main_frame.pack(fill=tk.BOTH, expand=True, pady=10)

        left_panel = tk.Frame(main_frame, bg=BG_PANEL, width=440)
        left_panel.pack(side=tk.LEFT, fill=tk.Y, padx=(0, 15))
        left_panel.pack_propagate(False)

        self.lbl_model_info = tk.Label(
            left_panel, text="", 
            font=("맑은 고딕", 11, "bold"), fg=ACCENT_YELLOW, bg=BG_PANEL, justify=tk.LEFT
        )
        self.lbl_model_info.pack(anchor="w", padx=20, pady=(15, 10))

        self.status_box = tk.Label(
            left_panel, text="READY", font=("Arial", 38, "bold"),
            fg="#adb5bd", bg="#2a2e37", height=3, relief="flat"
        )
        self.status_box.pack(fill=tk.X, padx=20, pady=6)

        self.lbl_last_scan = tk.Label(
            left_panel, text=f"{self.t('last_scan')}: -", font=("맑은 고딕", 9),
            fg=TEXT_MUTED, bg=BG_PANEL, anchor="w"
        )
        self.lbl_last_scan.pack(fill=tk.X, padx=20, pady=(10, 3))

        self.lbl_input_guide = tk.Label(left_panel, text=self.t("input_guide"), font=("맑은 고딕", 9),
                                        fg=TEXT_MUTED, bg=BG_PANEL, anchor="w")
        self.lbl_input_guide.pack(fill=tk.X, padx=20)

        self.scan_entry = tk.Entry(
            left_panel, font=("Consolas", 11), bg=BG_INPUT, fg="#ffffff",
            insertbackground="#ffffff", relief="flat", highlightthickness=1,
            highlightbackground="#343c4c", highlightcolor="#3b82f6"
        )
        self.scan_entry.pack(fill=tk.X, padx=20, pady=(4, 15), ipady=5)
        self.scan_entry.bind("<Return>", lambda e: self.process_scan(self.scan_entry.get()))
        self.scan_entry.bind("<KeyRelease>", self.on_entry_key_release)

        stats_frame = tk.Frame(left_panel, bg=BG_PANEL)
        stats_frame.pack(fill=tk.X, padx=20, pady=3)
        stats_frame.columnconfigure((0, 1, 2), weight=1)

        card_total = tk.Frame(stats_frame, bg="#1a1e26", pady=6)
        card_total.grid(row=0, column=0, padx=2, sticky="nsew")
        self.lbl_total_val = tk.Label(card_total, text="0", font=("Arial", 16, "bold"), fg=TEXT_COLOR, bg="#1a1e26")
        self.lbl_total_val.pack()
        tk.Label(card_total, text="TOTAL", font=("Arial", 8, "bold"), fg=TEXT_MUTED, bg="#1a1e26").pack()

        card_ok = tk.Frame(stats_frame, bg="#1a1e26", pady=6)
        card_ok.grid(row=0, column=1, padx=2, sticky="nsew")
        self.lbl_ok_val = tk.Label(card_ok, text="0", font=("Arial", 16, "bold"), fg="#28a745", bg="#1a1e26")
        self.lbl_ok_val.pack()
        tk.Label(card_ok, text="OK", font=("Arial", 8, "bold"), fg=TEXT_MUTED, bg="#1a1e26").pack()

        card_ng = tk.Frame(stats_frame, bg="#1a1e26", pady=6)
        card_ng.grid(row=0, column=2, padx=2, sticky="nsew")
        self.lbl_ng_val = tk.Label(card_ng, text="0", font=("Arial", 16, "bold"), fg="#dc3545", bg="#1a1e26")
        self.lbl_ng_val.pack()
        tk.Label(card_ng, text="NG", font=("Arial", 8, "bold"), fg=TEXT_MUTED, bg="#1a1e26").pack()

        btn_reset = tk.Button(
            left_panel, text=self.t("reset_btn"), command=self.open_reset_dialog,
            bg="#2c323d", fg="#ff8787", activebackground="#3d2729", activeforeground="#ff6b6b",
            relief="flat", font=("맑은 고딕", 9, "bold"), pady=4, cursor="hand2"
        )
        btn_reset.pack(fill=tk.X, padx=20, pady=(8, 4))
        self.btn_reset = btn_reset

        self.btn_manager = tk.Button(
            left_panel, text=self.t("manager_btn"), command=self.toggle_manager_mode,
            bg="#2c323d", fg="#adb5bd", activebackground="#303642", activeforeground="#ffffff",
            relief="flat", font=("맑은 고딕", 9, "bold"), pady=4, cursor="hand2"
        )
        self.btn_manager.pack(fill=tk.X, padx=20, pady=(0, 8))

        self.lbl_pallet_status = tk.Label(
            left_panel, text="현재 팔레트: - (0/12 박스)",
            font=("맑은 고딕", 9, "bold"), fg="#38bdf8", bg=BG_PANEL, anchor="w"
        )
        self.lbl_pallet_status.pack(fill=tk.X, padx=20, pady=(2, 3))

        self.lbl_pending_status = tk.Label(
            left_panel, text="", font=("맑은 고딕", 9), fg=TEXT_MUTED, bg=BG_PANEL, anchor="w"
        )
        self.lbl_pending_status.pack(fill=tk.X, padx=20, pady=(0, 5))

        right_panel = tk.Frame(main_frame, bg=BG_MAIN)
        right_panel.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True)

        self.lbl_right_header = tk.Label(
            right_panel, text="", font=("맑은 고딕", 10, "bold"), fg="#9aa0a6", bg=BG_MAIN, anchor="w"
        )
        self.lbl_right_header.pack(fill=tk.X, pady=(0, 6))

        columns = ("Pallet", "DAY", "TIME", "Label QR", "DMC", "JUDGMENT", "Content")
        self.tree = ttk.Treeview(right_panel, columns=columns, show="headings", style="Dark.Treeview")
        self.tree.tag_configure("ng_row", background="#3a1c1f", foreground="#ff6b6b")
        self.tree.tag_configure("pallet_row", background="#1e3a5f", foreground="#7dd3fc")

        self.tree.heading("Pallet", text=self.t("th_pallet"))
        self.tree.heading("DAY", text=self.t("th_day"))
        self.tree.heading("TIME", text=self.t("th_time"))
        self.tree.heading("Label QR", text=self.t("th_label"))
        self.tree.heading("DMC", text=self.t("th_dmc"))
        self.tree.heading("JUDGMENT", text=self.t("th_judgment"))
        self.tree.heading("Content", text=self.t("th_content"))

        self.tree.column("Pallet", width=170, anchor="w")
        self.tree.column("DAY", width=80, anchor="center")
        self.tree.column("TIME", width=70, anchor="center")
        self.tree.column("Label QR", width=220, anchor="w")
        self.tree.column("DMC", width=210, anchor="w")
        self.tree.column("JUDGMENT", width=75, anchor="center")
        self.tree.column("Content", width=95, anchor="center")

        tree_scroll = ttk.Scrollbar(right_panel, orient=tk.VERTICAL, command=self.tree.yview)
        self.tree.configure(yscroll=tree_scroll.set)

        self.tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        tree_scroll.pack(side=tk.RIGHT, fill=tk.Y)

    # ==========================================
    # Grouping 전용 탭 (DMC 제외, Pallet-Box Grouping 내역 전용)
    # ==========================================
    def build_grouping_tab(self):
        group_frame = tk.Frame(self.tab_grouping, bg=BG_MAIN)
        group_frame.pack(fill=tk.BOTH, expand=True, padx=15, pady=10)

        header_bar = tk.Frame(group_frame, bg=BG_PANEL, pady=8, padx=15)
        header_bar.pack(fill=tk.X, pady=(0, 10))

        self.lbl_grouping_header = tk.Label(
            header_bar, text="",
            font=("맑은 고딕", 10, "bold"), fg="#38bdf8", bg=BG_PANEL
        )
        self.lbl_grouping_header.pack(side=tk.LEFT)

        btn_refresh = tk.Button(
            header_bar, text="🔄 새로고침", command=self.refresh_grouping_tab,
            bg="#2b5278", fg="#ffffff", relief="flat", font=("맑은 고딕", 9, "bold"), padx=12, pady=2, cursor="hand2"
        )
        btn_refresh.pack(side=tk.RIGHT)

        cols = ("Pallet", "BoxSeq", "DAY", "TIME", "Label QR", "Qty", "JUDGMENT")
        self.tree_grouping = ttk.Treeview(group_frame, columns=cols, show="headings", style="Dark.Treeview")
        self.tree_grouping.tag_configure("pallet_start", background="#1a365d", foreground="#93c5fd")
        self.tree_grouping.tag_configure("pallet_done", background="#223042", foreground="#94a3b8")
        self.tree_grouping.tag_configure("ng_box", background="#3a1c1f", foreground="#ff6b6b")

        self.tree_grouping.heading("Pallet", text=self.t("th_pallet"))
        self.tree_grouping.heading("BoxSeq", text=self.t("th_box_seq"))
        self.tree_grouping.heading("DAY", text=self.t("th_day"))
        self.tree_grouping.heading("TIME", text=self.t("th_time"))
        self.tree_grouping.heading("Label QR", text=self.t("th_label"))
        self.tree_grouping.heading("Qty", text="수량 (Qty)")
        self.tree_grouping.heading("JUDGMENT", text=self.t("th_judgment"))

        self.tree_grouping.column("Pallet", width=220, anchor="w")
        self.tree_grouping.column("BoxSeq", width=85, anchor="center")
        self.tree_grouping.column("DAY", width=95, anchor="center")
        self.tree_grouping.column("TIME", width=85, anchor="center")
        self.tree_grouping.column("Label QR", width=340, anchor="w")
        self.tree_grouping.column("Qty", width=90, anchor="center")
        self.tree_grouping.column("JUDGMENT", width=85, anchor="center")

        scroll_g = ttk.Scrollbar(group_frame, orient=tk.VERTICAL, command=self.tree_grouping.yview)
        self.tree_grouping.configure(yscroll=scroll_g.set)

        self.tree_grouping.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scroll_g.pack(side=tk.RIGHT, fill=tk.Y)

    def refresh_grouping_tab(self):
        self.tree_grouping.delete(*self.tree_grouping.get_children())
        curr_model = self.current_model.get()
        files = self.get_all_model_files(curr_model)
        if not files:
            return

        def _loader():
            group_rows = []
            current_pallet_code = ""
            box_seq_tracker = 0

            for filepath in files:
                unhide_file(filepath)
                wb = openpyxl.load_workbook(filepath, data_only=True)
                ws = wb["스캔실적"] if "스캔실적" in wb.sheetnames else wb.active

                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or len(row) < 7:
                        continue

                    if len(row) >= 8:
                        p_val = str(row[0]).strip() if row[0] and str(row[0]).strip() != "-" else ""
                        lbl_val = str(row[1]).strip() if row[1] and str(row[1]).strip() != "-" else ""
                        box_time = str(row[2]).strip() if row[2] and str(row[2]).strip() != "-" else ""
                        seq_val = str(row[3]).strip() if row[3] else ""
                        desc_val = str(row[4]).strip() if row[4] else ""
                        res_val = str(row[6]).strip() if row[6] else "OK"
                    else:
                        p_val = ""
                        lbl_val = str(row[0]).strip() if row[0] and str(row[0]).strip() != "-" else ""
                        box_time = str(row[1]).strip() if row[1] and str(row[1]).strip() != "-" else ""
                        seq_val = str(row[2]).strip() if row[2] else ""
                        desc_val = str(row[3]).strip() if row[3] else ""
                        res_val = str(row[5]).strip() if row[5] else "OK"

                    if seq_val == "Final HEADER" and "Start" in desc_val:
                        current_pallet_code = p_val
                        box_seq_tracker = 0
                        t_parts = box_time.split()
                        d_str = t_parts[0] if len(t_parts) > 0 else ""
                        tm_str = t_parts[1] if len(t_parts) > 1 else ""
                        group_rows.append((p_val, "-", d_str, tm_str, "[Pallet Grouping Start]", "-", "START", "pallet_start"))
                        continue

                    if seq_val == "Final HEADER" and "Done" in desc_val:
                        t_parts = box_time.split()
                        d_str = t_parts[0] if len(t_parts) > 0 else ""
                        tm_str = t_parts[1] if len(t_parts) > 1 else ""
                        group_rows.append((p_val, "-", d_str, tm_str, "[Pallet Grouping Done]", f"{box_seq_tracker} 박스", "DONE", "pallet_done"))
                        box_seq_tracker = 0
                        continue

                    if seq_val == "HEADER" or "Group" in desc_val or "Box" in desc_val:
                        box_seq_tracker += 1
                        t_parts = box_time.split()
                        d_str = t_parts[0] if len(t_parts) > 0 else ""
                        tm_str = t_parts[1] if len(t_parts) > 1 else ""

                        qty_str = "10"
                        if "Done:" in desc_val:
                            try:
                                qty_str = desc_val.split("Done:")[1].split("pcs")[0].strip()
                            except Exception:
                                qty_str = "10"

                        tag = "ng_box" if res_val == "NG" else ""
                        effective_p = p_val if p_val else current_pallet_code
                        group_rows.append((effective_p, f"#{box_seq_tracker}", d_str, tm_str, lbl_val, f"{qty_str} pcs", res_val, tag))

                hide_file(filepath)

            def _populate():
                for r in reversed(group_rows):
                    tag = r[7]
                    self.tree_grouping.insert("", tk.END, values=r[:7], tags=(tag,) if tag else ())

            self.root.after(0, _populate)

        threading.Thread(target=_loader, daemon=True).start()

    def build_recode_tab(self):
        recode_frame = tk.Frame(self.tab_recode, bg=BG_MAIN)
        recode_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

        filter_bar = tk.Frame(recode_frame, bg=BG_PANEL, pady=8, padx=15)
        filter_bar.pack(fill=tk.X, pady=(0, 10))

        today_str = datetime.now().strftime("%Y-%m-%d")

        self.lbl_filter_day = tk.Label(filter_bar, text=self.t("filter_day"), font=("맑은 고딕", 9, "bold"), fg=TEXT_COLOR, bg=BG_PANEL)
        self.lbl_filter_day.pack(side=tk.LEFT)
        self.entry_start_day = tk.Entry(filter_bar, width=11, font=("맑은 고딕", 9), justify="center", bg=BG_INPUT, fg="#ffffff")
        self.entry_start_day.insert(0, today_str)
        self.entry_start_day.pack(side=tk.LEFT, padx=5)

        tk.Label(filter_bar, text="~", fg=TEXT_MUTED, bg=BG_PANEL).pack(side=tk.LEFT)
        self.entry_end_day = tk.Entry(filter_bar, width=11, font=("맑은 고딕", 9), justify="center", bg=BG_INPUT, fg="#ffffff")
        self.entry_end_day.insert(0, today_str)
        self.entry_end_day.pack(side=tk.LEFT, padx=5)

        self.lbl_filter_time = tk.Label(filter_bar, text=self.t("filter_time"), font=("맑은 고딕", 9, "bold"), fg=TEXT_COLOR, bg=BG_PANEL)
        self.lbl_filter_time.pack(side=tk.LEFT, padx=(15, 0))
        self.entry_start_time = tk.Entry(filter_bar, width=9, font=("맑은 고딕", 9), justify="center", bg=BG_INPUT, fg="#ffffff")
        self.entry_start_time.insert(0, "00:00:00")
        self.entry_start_time.pack(side=tk.LEFT, padx=5)

        tk.Label(filter_bar, text="~", fg=TEXT_MUTED, bg=BG_PANEL).pack(side=tk.LEFT)
        self.entry_end_time = tk.Entry(filter_bar, width=9, font=("맑은 고딕", 9), justify="center", bg=BG_INPUT, fg="#ffffff")
        self.entry_end_time.insert(0, "23:59:59")
        self.entry_end_time.pack(side=tk.LEFT, padx=5)

        self.btn_search = tk.Button(
            filter_bar, text=self.t("search_btn"), command=self.apply_recode_filter,
            bg="#2b5278", fg="#ffffff", relief="flat", font=("맑은 고딕", 9, "bold"), padx=12, pady=2, cursor="hand2"
        )
        self.btn_search.pack(side=tk.LEFT, padx=15)

        self.btn_save = tk.Button(
            filter_bar, text=self.t("save_btn"), command=self.save_recode_to_excel,
            bg="#198754", fg="#ffffff", relief="flat", font=("맑은 고딕", 9, "bold"), padx=12, pady=2, cursor="hand2"
        )
        self.btn_save.pack(side=tk.RIGHT)

        cols = ("Pallet", "DAY", "TIME", "Label QR", "DMC", "JUDGMENT", "Content")
        self.tree_recode = ttk.Treeview(recode_frame, columns=cols, show="headings", style="Dark.Treeview")
        self.tree_recode.tag_configure("ng_row", background="#3a1c1f", foreground="#ff6b6b")
        self.tree_recode.tag_configure("pallet_row", background="#1e3a5f", foreground="#7dd3fc")

        self.tree_recode.heading("Pallet", text=self.t("th_pallet"))
        self.tree_recode.heading("DAY", text=self.t("th_day"))
        self.tree_recode.heading("TIME", text=self.t("th_time"))
        self.tree_recode.heading("Label QR", text=self.t("th_label"))
        self.tree_recode.heading("DMC", text=self.t("th_dmc"))
        self.tree_recode.heading("JUDGMENT", text=self.t("th_judgment"))
        self.tree_recode.heading("Content", text=self.t("th_content"))

        self.tree_recode.column("Pallet", width=170, anchor="w")
        self.tree_recode.column("DAY", width=80, anchor="center")
        self.tree_recode.column("TIME", width=70, anchor="center")
        self.tree_recode.column("Label QR", width=220, anchor="w")
        self.tree_recode.column("DMC", width=210, anchor="w")
        self.tree_recode.column("JUDGMENT", width=75, anchor="center")
        self.tree_recode.column("Content", width=95, anchor="center")

        scroll_r = ttk.Scrollbar(recode_frame, orient=tk.VERTICAL, command=self.tree_recode.yview)
        self.tree_recode.configure(yscroll=scroll_r.set)

        self.tree_recode.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scroll_r.pack(side=tk.RIGHT, fill=tk.Y)

    def on_language_changed(self, event=None):
        self.notebook.tab(0, text=self.t("tab_scan"))
        self.notebook.tab(1, text=self.t("tab_grouping"))
        self.notebook.tab(2, text=self.t("tab_recode"))

        self.btn_pw.config(text=self.t("pw_setting"))
        self.btn_reset.config(text=self.t("reset_btn"))
        if self.is_manager_mode:
            self.btn_manager.config(text=self.t("manager_btn_on"))
        else:
            self.btn_manager.config(text=self.t("manager_btn"))

        self.lbl_input_guide.config(text=self.t("input_guide"))
        model = self.current_model.get()
        target_code = MODEL_CONFIG[model]
        self.lbl_model_info.config(text=f"{self.t('model_label')} {model}\n{self.t('code_label')} {target_code}")
        self.lbl_right_header.config(text=self.t("record_header", model=model))
        self.lbl_grouping_header.config(text=self.t("grouping_header", model=model))
        self.lbl_pending_status.config(text=self.t("pending_status", count=len(self.pending_items)))
        self.update_pallet_status_ui()

        for tree_obj in (self.tree, self.tree_recode):
            tree_obj.heading("Pallet", text=self.t("th_pallet"))
            tree_obj.heading("DAY", text=self.t("th_day"))
            tree_obj.heading("TIME", text=self.t("th_time"))
            tree_obj.heading("Label QR", text=self.t("th_label"))
            tree_obj.heading("DMC", text=self.t("th_dmc"))
            tree_obj.heading("JUDGMENT", text=self.t("th_judgment"))
            tree_obj.heading("Content", text=self.t("th_content"))

        self.tree_grouping.heading("Pallet", text=self.t("th_pallet"))
        self.tree_grouping.heading("BoxSeq", text=self.t("th_box_seq"))
        self.tree_grouping.heading("DAY", text=self.t("th_day"))
        self.tree_grouping.heading("TIME", text=self.t("th_time"))
        self.tree_grouping.heading("Label QR", text=self.t("th_label"))
        self.tree_grouping.heading("JUDGMENT", text=self.t("th_judgment"))

        self.lbl_filter_day.config(text=self.t("filter_day"))
        self.lbl_filter_time.config(text=self.t("filter_time"))
        self.btn_search.config(text=self.t("search_btn"))
        self.btn_save.config(text=self.t("save_btn"))

        self.scan_entry.focus_set()

    def update_pallet_status_ui(self):
        m = self.current_model.get()
        info = self.pallet_state.get(m, {"current_pallet": "", "box_count": 0})
        p_name = info["current_pallet"] if info["current_pallet"] else "-"
        b_cnt = info["box_count"]
        self.lbl_pallet_status.config(text=self.t("pallet_status", pallet=p_name, boxes=b_cnt, max_b=MAX_BOXES_PER_PALLET))

    def on_entry_key_release(self, event):
        if event.keysym in ("Return", "KP_Enter"):
            if self.auto_submit_timer:
                self.root.after_cancel(self.auto_submit_timer)
                self.auto_submit_timer = None
            return

        text = self.scan_entry.get().strip()
        if len(text) >= 10:
            if self.auto_submit_timer:
                self.root.after_cancel(self.auto_submit_timer)
            self.auto_submit_timer = self.root.after(150, self._trigger_auto_submit_entry)

    def _trigger_auto_submit_entry(self):
        text = self.scan_entry.get().strip()
        if text:
            self.process_scan(text)

    def setup_global_key_listener(self):
        def _on_key_press(event):
            focused = self.root.focus_get()
            if isinstance(focused, tk.Entry) and focused != self.scan_entry:
                return

            if event.keysym in ("Return", "KP_Enter"):
                if self.auto_submit_timer:
                    self.root.after_cancel(self.auto_submit_timer)
                    self.auto_submit_timer = None

                if self.global_scan_buffer:
                    scanned_text = "".join(self.global_scan_buffer).strip()
                    self.global_scan_buffer.clear()
                    self.scan_entry.delete(0, tk.END)
                    if scanned_text:
                        self.process_scan(scanned_text)
            elif event.char and event.char.isprintable():
                self.global_scan_buffer.append(event.char)
                if len(self.global_scan_buffer) >= 10:
                    if self.auto_submit_timer:
                        self.root.after_cancel(self.auto_submit_timer)
                    self.auto_submit_timer = self.root.after(150, self._trigger_auto_submit_global)

        self.root.bind_all("<Key>", _on_key_press)

    def _trigger_auto_submit_global(self):
        if self.global_scan_buffer:
            scanned_text = "".join(self.global_scan_buffer).strip()
            self.global_scan_buffer.clear()
            self.scan_entry.delete(0, tk.END)
            if scanned_text:
                self.process_scan(scanned_text)

    def set_status(self, text, fg_color, bg_color):
        if len(text) <= 2:
            font_size = 46
        elif len(text) <= 7:
            font_size = 36
        elif len(text) <= 12:
            font_size = 28
        else:
            font_size = 22

        self.status_box.config(text=text, fg=fg_color, bg=bg_color, font=("Arial", font_size, "bold"))

    def center_popup(self, dialog, width, height):
        self.root.update_idletasks()
        rx = self.root.winfo_x()
        ry = self.root.winfo_y()
        rw = self.root.winfo_width()
        rh = self.root.winfo_height()

        x = rx + (rw - width) // 2
        y = ry + (rh - height) // 2
        dialog.geometry(f"{width}x{height}+{max(0, x)}+{max(0, y)}")

    def toggle_manager_mode(self):
        if self.active_popup or self.pallet_wait_popup:
            return

        win = tk.Toplevel(self.root)
        win.configure(bg=BG_PANEL)
        win.transient(self.root)
        win.grab_set()

        self.center_popup(win, 360, 210)

        target_state = not self.is_manager_mode
        if target_state:
            win.title(self.t("manager_btn") + " ON")
            msg_text = "MANAGER MODE [ON]\n" + ("관리자 비밀번호를 입력하세요." if self.current_lang.get()=="한국어" else "Enter Admin Password.")
        else:
            win.title(self.t("manager_btn") + " OFF")
            msg_text = "MANAGER MODE [OFF]\n" + ("관리자 비밀번호를 입력하세요." if self.current_lang.get()=="한국어" else "Enter Admin Password.")

        tk.Label(win, text=msg_text, font=("맑은 고딕", 10, "bold"), fg=TEXT_COLOR, bg=BG_PANEL).pack(pady=(15, 8))

        pw_entry = tk.Entry(win, show="*", font=("Arial", 14), justify="center", bg=BG_INPUT, fg="#ffffff")
        pw_entry.pack(pady=5)
        pw_entry.focus_set()

        lbl_err = tk.Label(win, text="", font=("맑은 고딕", 9), fg="#ff6b6b", bg=BG_PANEL)
        lbl_err.pack()

        def verify(event=None):
            if pw_entry.get() == self.admin_password:
                self.is_manager_mode = target_state
                if self.is_manager_mode:
                    self.btn_manager.config(bg="#28a745", fg="#ffffff", text=self.t("manager_btn_on"))
                else:
                    self.btn_manager.config(bg="#2c323d", fg="#adb5bd", text=self.t("manager_btn"))
                win.destroy()
                self.scan_entry.focus_set()
            else:
                lbl_err.config(text=self.t("pw_err"))
                pw_entry.delete(0, tk.END)

        pw_entry.bind("<Return>", verify)
        tk.Button(win, text=self.t("unlock_btn"), command=verify, bg="#28a745" if target_state else "#dc3545", fg="#ffffff",
                  relief="flat", font=("맑은 고딕", 10, "bold"), padx=15, pady=3).pack(pady=10)

    def auto_turn_off_manager_mode(self):
        self.is_manager_mode = False
        self.btn_manager.config(bg="#2c323d", fg="#adb5bd", text=self.t("manager_btn"))

    def open_sorting_popup(self, dmc_code):
        dialog = tk.Toplevel(self.root)
        dialog.title(self.t("sorting_title"))
        dialog.resizable(False, False)
        dialog.configure(bg="#3a1c1f")

        dialog.transient(self.root)
        dialog.grab_set()
        dialog.protocol("WM_DELETE_WINDOW", lambda: None)

        self.center_popup(dialog, 520, 300)
        self.active_popup = dialog

        self.play_alarm_sound()

        msg = self.t("sorting_msg", code=dmc_code)
        tk.Label(dialog, text=msg, font=("맑은 고딕", 11, "bold"), bg="#3a1c1f", fg="#ff6b6b", justify=tk.LEFT).pack(pady=25, padx=20)

        def close_dialog(event=None):
            dialog.grab_release()
            dialog.destroy()
            self.active_popup = None
            self.set_status("READY", "#adb5bd", "#2a2e37")
            self.scan_entry.focus_set()

        dialog.bind("<Return>", close_dialog)
        dialog.bind("<KP_Enter>", close_dialog)

        btn = tk.Button(dialog, text=self.t("confirm_btn"), command=close_dialog,
                        font=("맑은 고딕", 11, "bold"), bg="#dc3545", fg="#ffffff",
                        relief="flat", padx=20, pady=6, cursor="hand2")
        btn.pack(pady=10)
        btn.focus_set()

    def open_pallet_wait_popup(self):
        if self.pallet_wait_popup:
            return

        dialog = tk.Toplevel(self.root)
        dialog.title(self.t("pallet_popup_title"))
        dialog.resizable(False, False)
        dialog.configure(bg="#1e293b")

        dialog.transient(self.root)
        dialog.grab_set()
        dialog.protocol("WM_DELETE_WINDOW", lambda: None)

        self.center_popup(dialog, 520, 240)
        self.pallet_wait_popup = dialog

        tk.Label(dialog, text="[12 BOX PACKING COMPLETE]", font=("맑은 고딕", 12, "bold"), fg="#38bdf8", bg="#1e293b").pack(pady=(25, 8))
        tk.Label(dialog, text=self.t("pallet_popup_msg"), font=("맑은 고딕", 13, "bold"), fg="#f8fafc", bg="#1e293b").pack(pady=10)
        tk.Label(dialog, text="* 새 Pallet QR을 스캔하면 자동으로 해제됩니다.", font=("맑은 고딕", 9), fg="#94a3b8", bg="#1e293b").pack(pady=5)

    def close_pallet_wait_popup(self):
        if self.pallet_wait_popup:
            try:
                self.pallet_wait_popup.grab_release()
                self.pallet_wait_popup.destroy()
            except Exception:
                pass
            self.pallet_wait_popup = None

    # ==========================================
    # 4. 모델 변경 및 과거 기록 복원 로더
    # ==========================================
    def on_model_changed(self, event=None):
        self.model_session_id += 1
        current_session = self.model_session_id

        model = self.current_model.get()
        target_code = MODEL_CONFIG[model]

        self.lbl_model_info.config(text=f"{self.t('model_label')} {model}\n{self.t('code_label')} {target_code}")
        self.lbl_right_header.config(text=self.t("record_header", model=model))
        self.lbl_grouping_header.config(text=self.t("grouping_header", model=model))
        
        self.pending_items.clear()
        self.pending_tree_ids.clear()
        self.lbl_pending_status.config(text=self.t("pending_status", count=0))
        self.update_pallet_status_ui()

        self.update_stat_cards()
        self.load_history_from_excel(model, target_code, current_session)
        self.refresh_grouping_tab()
        self.scan_entry.focus_set()

    def get_all_model_files(self, model_name):
        safe_model = model_name.replace('-', '_')
        pattern = os.path.join(BASE_DIR, f"Y*_*Q_{safe_model}.xlsx")
        files = glob.glob(pattern)
        old_file = os.path.join(BASE_DIR, f"{model_name}.xlsx")
        if os.path.exists(old_file) and old_file not in files:
            files.append(old_file)
        files.sort()
        return files

    def load_history_from_excel(self, model_name, target_code, session_id):
        self.tree.delete(*self.tree.get_children())
        files = self.get_all_model_files(model_name)
        if not files:
            return

        def _loader():
            try:
                rows_to_insert = []
                target_upper = target_code.upper()
                last_label = ""
                loaded_pending = []
                self.sorting_list_by_model[model_name].clear()

                for filepath in files:
                    unhide_file(filepath)
                    wb = openpyxl.load_workbook(filepath, data_only=True)

                    if "sorting" in wb.sheetnames:
                        ws_sort = wb["sorting"]
                        for r in range(2, 2001):
                            cell_v = ws_sort.cell(row=r, column=2).value
                            if cell_v:
                                self.sorting_list_by_model[model_name].add(str(cell_v).strip().upper())

                    if "스캔실적" in wb.sheetnames:
                        ws = wb["스캔실적"]
                    else:
                        sheets = [s for s in wb.worksheets if s.title != "sorting"]
                        ws = sheets[0] if sheets else wb.active

                    for row in ws.iter_rows(min_row=2, values_only=True):
                        if not row or len(row) < 7:
                            continue
                        
                        if len(row) >= 8:
                            pallet_val = str(row[0]).strip().upper() if row[0] and str(row[0]).strip() != "-" else ""
                            lbl_val = row[1]
                            box_time = row[2]
                            seq_val = row[3]
                            dmc_val = row[4]
                            dmc_time = row[5]
                            res = row[6]
                            content = row[7] if len(row) >= 8 and row[7] else ""
                        else:
                            pallet_val = ""
                            lbl_val = row[0]
                            box_time = row[1]
                            seq_val = row[2]
                            dmc_val = row[3]
                            dmc_time = row[4]
                            res = row[5]
                            content = row[6] if len(row) >= 7 and row[6] else ""

                        lbl_str = str(lbl_val).strip() if lbl_val and str(lbl_val).strip() != "-" else ""
                        dmc_str = str(dmc_val).strip() if dmc_val and str(dmc_val).strip() != "-" else ""
                        res_str = str(res).strip() if res else "OK"
                        content_str = str(content).strip() if content else ""

                        if seq_val == "Final HEADER":
                            rows_to_insert.append((pallet_val, "", "", "", dmc_str, res_str, content_str))
                            continue

                        if seq_val == "HEADER" or dmc_str.startswith("[박스 묶음 완료") or "Group" in dmc_str or "Pakiet" in dmc_str:
                            t_parts = str(box_time).split()
                            day_val = t_parts[0] if len(t_parts) > 0 else ""
                            time_val = t_parts[1] if len(t_parts) > 1 else ""
                            rows_to_insert.append((pallet_val, day_val, time_val, lbl_str, dmc_str, res_str, content_str))
                            if lbl_str:
                                self.scanned_label_by_model[model_name].add(lbl_str)
                            last_label = ""
                            loaded_pending.clear()
                            continue

                        ts_str = str(dmc_time if dmc_time and str(dmc_time).strip() != "-" else box_time)
                        if not ts_str or ts_str == "-":
                            continue

                        t_parts = ts_str.split()
                        day_val = t_parts[0] if len(t_parts) > 0 else ""
                        time_val = t_parts[1] if len(t_parts) > 1 else ""

                        clean_dmc = dmc_str.replace("[중복스캔] ", "").replace("[중복 스캔] ", "").upper()
                        if not clean_dmc.startswith(target_upper):
                            continue

                        if lbl_str:
                            last_label = lbl_str
                            final_label = lbl_str
                        else:
                            final_label = ""
                            if res_str == "OK":
                                loaded_pending.append({
                                    "day": day_val, "time": time_val, "code": dmc_str, "result": "OK"
                                })

                        if final_label and final_label.upper().startswith(target_upper):
                            self.scanned_label_by_model[model_name].add(final_label)
                        if clean_dmc and clean_dmc.startswith(target_upper):
                            self.scanned_history_by_model[model_name].add(clean_dmc)

                        rows_to_insert.append((pallet_val, day_val, time_val, final_label, dmc_str, res_str, content_str))

                    hide_file(filepath)

                if session_id != self.model_session_id:
                    return

                def _populate():
                    if session_id == self.model_session_id:
                        for r in reversed(rows_to_insert):
                            tag = ""
                            if r[5] == "NG":
                                tag = "ng_row"
                            elif "Pallet" in str(r[4]):
                                tag = "pallet_row"
                            self.tree.insert("", tk.END, values=r, tags=(tag,) if tag else ())
                        
                        if loaded_pending:
                            self.pending_items = list(loaded_pending)
                            self.lbl_pending_status.config(text=self.t("pending_status", count=len(self.pending_items)))

                self.root.after(0, _populate)

            except Exception:
                pass

        threading.Thread(target=_loader, daemon=True).start()

    # ==========================================
    # 5. 스캔 판정 로직 (Pallet QR 모델 검증 탑재)
    # ==========================================
    def is_pallet_qr(self, code):
        c = code.strip().upper()
        if ';' in c:
            return False
        if c.startswith("PALLET") or c.startswith("KR02") or " " in c or "\t" in c:
            return True
        return False

    def process_scan(self, raw_code):
        if self.auto_submit_timer:
            self.root.after_cancel(self.auto_submit_timer)
            self.auto_submit_timer = None

        raw_code = raw_code.strip()
        self.scan_entry.delete(0, tk.END)
        self.global_scan_buffer.clear()

        if not raw_code:
            return

        if self.active_popup:
            return

        now = datetime.now()
        day_str = now.strftime("%Y-%m-%d")
        time_str = now.strftime("%H:%M:%S")
        timestamp_full = f"{day_str} {time_str}"

        curr_model = self.current_model.get()
        target_code = MODEL_CONFIG[curr_model].upper()

        # ----------------------------------------------------
        # Pallet QR 스캔 처리
        # ----------------------------------------------------
        if self.is_pallet_qr(raw_code):
            upper_pallet_code = raw_code.upper()  # 1. 무조건 대문자 변환

            # 2. Pallet QR 내부 ITEM 코드 일치성 검증 (모델 불일치 차단)
            if target_code not in upper_pallet_code:
                self.set_status("Pallet NG", "#dc3545", "#3a1c1f")
                self.open_lock_popup(
                    title_text=self.t("ng_pallet_model_title"),
                    msg=self.t("ng_pallet_model_msg", model=curr_model, target=target_code, code=upper_pallet_code),
                    header_bg="#2d1d20", header_fg="#f87171"
                )
                return

            if len(self.pending_items) > 0:
                self.set_status("Pallet NG", "#dc3545", "#3a1c1f")
                self.open_lock_popup(
                    title_text=self.t("ng_pallet_mid_title"),
                    msg=self.t("ng_pallet_mid_msg", count=len(self.pending_items)),
                    header_bg="#2d1d20", header_fg="#f87171"
                )
                return

            prev_pallet = self.pallet_state[curr_model]["current_pallet"]
            
            if prev_pallet:
                self.direct_append_pallet_header(curr_model, prev_pallet, "Final HEADER", "[Pallet Grouping Done]", timestamp_full)
                self.tree.insert("", 0, values=(prev_pallet, "", "", "", "[Pallet Grouping Done]", "OK", ""), tags=("pallet_row",))

            self.pallet_state[curr_model]["current_pallet"] = upper_pallet_code
            self.pallet_state[curr_model]["box_count"] = 0
            self.save_pallet_state()
            self.update_pallet_status_ui()

            self.direct_append_pallet_header(curr_model, upper_pallet_code, "Final HEADER", "[Pallet Grouping Start]", timestamp_full)
            self.tree.insert("", 0, values=(upper_pallet_code, "", "", "", "[Pallet Grouping Start]", "OK", ""), tags=("pallet_row",))

            self.close_pallet_wait_popup()
            self.refresh_grouping_tab()
            self.set_status("OK", "#28a745", "#193322")
            return

        # ----------------------------------------------------
        # 일반 바코드 (단품 및 Label QR) 처리
        # ----------------------------------------------------
        current_time = time.time()
        if raw_code == self.last_scanned_code and (current_time - self.last_scanned_time) < 2.0:
            return

        self.last_scanned_code = raw_code
        self.last_scanned_time = current_time

        scanned_prefix = raw_code[:10].upper()
        clean_upper_code = raw_code.upper()
        is_label_qr = (raw_code.count(';') >= 3)
        self.lbl_last_scan.config(text=f"{self.t('last_scan')}: {raw_code}")

        # [검증 1] 모델 코드 불일치 NG
        if scanned_prefix != target_code:
            self.set_status("NG", "#dc3545", "#3a1c1f")
            self.open_lock_popup(
                title_text=self.t("ng_model_title"),
                msg=self.t("ng_model_msg", model=curr_model, target=target_code, code=raw_code[:12]),
                header_bg="#2d1d20", header_fg="#f87171"
            )
            return

        # [검증 2] Sorting 필요 제품 체크 (C열 OK 마킹 및 팝업)
        if not is_label_qr and clean_upper_code in self.sorting_list_by_model[curr_model]:
            self.set_status("SORTING", "#f59f00", "#3d2716")
            self.direct_mark_sorting_ok(curr_model, raw_code)
            self.open_sorting_popup(raw_code)
            return

        # [검증 3] Label QR 전용 검증
        if is_label_qr:
            curr_box_cnt = self.pallet_state[curr_model]["box_count"]

            if curr_box_cnt >= MAX_BOXES_PER_PALLET:
                self.set_status("Pallet NG", "#dc3545", "#3a1c1f")
                self.open_lock_popup(
                    title_text=self.t("ng_pallet_limit_title"),
                    msg=self.t("ng_pallet_limit_msg"),
                    header_bg="#2d1d20", header_fg="#f87171"
                )
                return

            if raw_code in self.scanned_label_by_model[curr_model]:
                self.set_status("Label QR NG", "#dc3545", "#3a1c1f")
                self.open_lock_popup(
                    title_text=self.t("ng_label_dup_title"),
                    msg=self.t("ng_label_dup_msg", code=raw_code[:35]),
                    header_bg="#2d1d20", header_fg="#f87171"
                )
                return

            tokens = raw_code.split(';')
            expected_qty = None
            if len(tokens) >= 2 and tokens[1].isdigit():
                expected_qty = int(tokens[1])

            current_scanned_qty = len(self.pending_items)

            if expected_qty is not None and expected_qty != current_scanned_qty:
                self.set_status("Grouping NG", "#dc3545", "#3a1c1f")
                self.open_lock_popup(
                    title_text=self.t("ng_group_title"),
                    msg=self.t("ng_group_msg", expected=expected_qty, current=current_scanned_qty),
                    header_bg="#2d1d20", header_fg="#f87171"
                )
                return

        # [검증 4] 단품 QR 전용 체크
        if not is_label_qr:
            if len(self.pending_items) >= MAX_ITEMS_PER_BOX:
                self.set_status("NG", "#dc3545", "#3a1c1f")
                self.open_lock_popup(
                    title_text=self.t("ng_limit_title"),
                    msg=self.t("ng_limit_msg", max_cnt=MAX_ITEMS_PER_BOX),
                    header_bg="#2d1d20", header_fg="#f87171"
                )
                return

            is_already_scanned = (clean_upper_code in self.scanned_history_by_model[curr_model])

            if self.is_manager_mode:
                if not is_already_scanned:
                    self.set_status("NG", "#dc3545", "#3a1c1f")
                    self.open_lock_popup(
                        title_text=self.t("ng_mgr_err_title"),
                        msg=self.t("ng_mgr_err_msg", code=raw_code),
                        header_bg="#2d1d20", header_fg="#f87171"
                    )
                    return
                else:
                    self.set_status("OK", "#28a745", "#193322")
                    cur_p = self.pallet_state[curr_model]["current_pallet"]
                    item_data = {
                        "pallet": cur_p,
                        "day": day_str,
                        "time": time_str,
                        "code": raw_code,
                        "result": "OK"
                    }
                    self.tree.insert("", 0, values=(cur_p, day_str, time_str, "", raw_code, "OK", ""))
                    self.direct_append_single_item(curr_model, item_data)
                    self.auto_turn_off_manager_mode()
                    return

            else:
                if is_already_scanned:
                    self.model_counts[curr_model]["ng"] += 1
                    self.model_counts[curr_model]["total"] += 1
                    self.save_model_counts()
                    self.update_stat_cards()

                    self.set_status("QR NG", "#fd7e14", "#3d2716")

                    cur_p = self.pallet_state[curr_model]["current_pallet"]
                    dup_text = self.t("dup_scan_tag")
                    self.tree.insert("", 0, values=(cur_p, day_str, time_str, "-", raw_code, "NG", dup_text), tags=("ng_row",))

                    matched_label_qr = None
                    for item_id in self.tree.get_children():
                        vals = list(self.tree.item(item_id, "values"))
                        if not vals:
                            continue
                        if vals[4].strip().upper() == clean_upper_code and vals[5] != "NG":
                            matched_label_qr = vals[3]
                            vals[5] = "NG"
                            vals[6] = dup_text
                            self.tree.item(item_id, values=vals, tags=("ng_row",))
                            break

                    if matched_label_qr and matched_label_qr != "-":
                        for item_id in self.tree.get_children():
                            vals = list(self.tree.item(item_id, "values"))
                            if vals and vals[3] == matched_label_qr and ("[" in str(vals[4])):
                                vals[5] = "NG"
                                self.tree.item(item_id, values=vals, tags=("ng_row",))
                                break

                    self.direct_handle_dmc_duplicate(curr_model, raw_code, day_str, time_str, matched_label_qr, dup_text, cur_p)

                    self.open_lock_popup(
                        title_text=self.t("ng_dup_title"),
                        msg=self.t("ng_dup_msg", code=raw_code),
                        header_bg="#352316", header_fg="#fb923c"
                    )
                    return

        self.set_status("OK", "#28a745", "#193322")
        cur_pallet = self.pallet_state[curr_model]["current_pallet"]

        if not is_label_qr:
            self.model_counts[curr_model]["ok"] += 1
            self.model_counts[curr_model]["total"] += 1
            self.save_model_counts()
            self.update_stat_cards()

            self.scanned_history_by_model[curr_model].add(clean_upper_code)
            item_data = {
                "pallet": cur_pallet,
                "day": day_str,
                "time": time_str,
                "code": raw_code,
                "result": "OK"
            }
            self.pending_items.append(item_data)
            
            item_id = self.tree.insert("", 0, values=(cur_pallet, day_str, time_str, "", raw_code, "OK", ""))
            self.pending_tree_ids.append(item_id)
            self.lbl_pending_status.config(text=self.t("pending_status", count=len(self.pending_items)))

            self.direct_append_single_item(curr_model, item_data)

        else:
            self.scanned_label_by_model[curr_model].add(raw_code)

            for t_id in self.pending_tree_ids:
                curr_vals = self.tree.item(t_id, "values")
                if curr_vals:
                    self.tree.item(t_id, values=(cur_pallet, curr_vals[1], curr_vals[2], raw_code, curr_vals[4], curr_vals[5], curr_vals[6]))

            items_to_bundle = list(self.pending_items)
            bundle_count = len(items_to_bundle)
            header_text = self.t("box_complete", count=bundle_count)

            self.tree.insert("", 0, values=(cur_pallet, day_str, time_str, raw_code, header_text, "OK", ""))

            self.pending_items.clear()
            self.pending_tree_ids.clear()
            self.lbl_pending_status.config(text=self.t("pending_status", count=0))

            self.pallet_state[curr_model]["box_count"] += 1
            self.save_pallet_state()
            self.update_pallet_status_ui()

            self.root.update_idletasks()

            self.direct_finalize_excel_group(curr_model, cur_pallet, raw_code, timestamp_full, items_to_bundle, header_text)
            self.refresh_grouping_tab()

            if self.pallet_state[curr_model]["box_count"] >= MAX_BOXES_PER_PALLET:
                self.open_pallet_wait_popup()

        self.scan_entry.focus_set()

    # ==========================================
    # 6. 엑셀 8개 열(Pallet 열 포함) 안전 I/O
    # ==========================================
    def open_or_init_workbook(self, filepath):
        unhide_file(filepath)
        if os.path.exists(filepath):
            try:
                wb = openpyxl.load_workbook(filepath)
                if "스캔실적" in wb.sheetnames:
                    ws = wb["스캔실적"]
                else:
                    sheets = [s for s in wb.worksheets if s.title != "sorting"]
                    ws = sheets[0] if sheets else wb.create_sheet(title="스캔실적", index=0)
                
                if ws.cell(row=1, column=1).value != "Pallet Label QR":
                    ws.insert_cols(1)
                    ws.cell(row=1, column=1, value="Pallet Label QR")
                    header_fill = PatternFill(start_color="1F242D", end_color="1F242D", fill_type="solid")
                    header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
                    ws.cell(row=1, column=1).fill = header_fill
                    ws.cell(row=1, column=1).font = header_font
                    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center", vertical="center")
                    ws.column_dimensions['A'].width = 30
                return wb, ws
            except Exception:
                pass

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "스캔실적"

        headers = ["Pallet Label QR", "Label QR (Box/Lot)", "Label 스캔일시", "단품 순번", "단품 DMC", "단품 스캔일시", "판정", "Content"]
        ws.append(headers)

        header_fill = PatternFill(start_color="1F242D", end_color="1F242D", fill_type="solid")
        header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")

        for col in range(1, 9):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 46
        ws.column_dimensions['C'].width = 20
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 34
        ws.column_dimensions['F'].width = 20
        ws.column_dimensions['G'].width = 14
        ws.column_dimensions['H'].width = 16

        ws_sort = wb.create_sheet(title="sorting")
        ws_sort.cell(row=1, column=2, value="Sorting 대상 DMC Code")
        ws_sort.cell(row=1, column=3, value="판정")
        ws_sort.column_dimensions['B'].width = 35
        ws_sort.column_dimensions['C'].width = 12

        return wb, ws

    def direct_append_pallet_header(self, model_name, pallet_code, seq_val, text_val, timestamp_full):
        with self.file_lock:
            try:
                filename = get_quarter_filename(model_name)
                filepath = os.path.join(BASE_DIR, filename)
                wb, ws = self.open_or_init_workbook(filepath)

                thin_border = Border(
                    left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                    top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
                )
                start_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
                done_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

                fill_to_use = start_fill if "Start" in text_val else done_fill
                row_data = [pallet_code, "-", "-", seq_val, text_val, timestamp_full, "OK", ""]
                ws.append(row_data)
                h_idx = ws.max_row
                for col in range(1, 9):
                    c = ws.cell(row=h_idx, column=col)
                    c.border = thin_border
                    c.fill = fill_to_use
                    c.alignment = Alignment(horizontal="center" if col in [3, 4, 6, 7, 8] else "left", vertical="center")

                wb.save(filepath)
                hide_file(filepath)
            except Exception as e:
                print(f"[팔레트 헤더 저장 실패]: {e}")

    def direct_mark_sorting_ok(self, model_name, raw_code):
        with self.file_lock:
            try:
                clean_target = raw_code.strip().upper()
                files = self.get_all_model_files(model_name)
                current_quarter_file = os.path.join(BASE_DIR, get_quarter_filename(model_name))
                if current_quarter_file not in files:
                    files.append(current_quarter_file)

                ok_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
                ok_font = Font(name="맑은 고딕", size=10, bold=True, color="006100")
                thin_border = Border(
                    left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                    top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
                )

                for f_path in files:
                    if not os.path.exists(f_path):
                        continue
                    unhide_file(f_path)
                    wb = openpyxl.load_workbook(f_path)
                    if "sorting" in wb.sheetnames:
                        ws_sort = wb["sorting"]
                        modified = False
                        for r in range(2, 2001):
                            val = ws_sort.cell(row=r, column=2).value
                            if val and str(val).strip().upper() == clean_target:
                                c_cell = ws_sort.cell(row=r, column=3)
                                c_cell.value = "OK"
                                c_cell.fill = ok_fill
                                c_cell.font = ok_font
                                c_cell.border = thin_border
                                c_cell.alignment = Alignment(horizontal="center", vertical="center")
                                modified = True
                        if modified:
                            wb.save(f_path)
                    hide_file(f_path)
            except Exception as e:
                print(f"[Sorting OK 마킹 오류]: {e}")

    def direct_append_single_item(self, model_name, item):
        with self.file_lock:
            try:
                filename = get_quarter_filename(model_name)
                filepath = os.path.join(BASE_DIR, filename)
                wb, ws = self.open_or_init_workbook(filepath)

                thin_border = Border(
                    left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                    top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
                )
                ok_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

                start_row = ws.max_row + 1
                ts_full = f"{item['day']} {item['time']}"
                row_data = [item.get("pallet", ""), "-", "-", "-", item["code"], ts_full, item["result"], ""]
                ws.append(row_data)

                for col in range(1, 9):
                    c = ws.cell(row=start_row, column=col)
                    c.border = thin_border
                    c.alignment = Alignment(horizontal="center" if col in [3, 4, 6, 7, 8] else "left", vertical="center")
                    if col == 7:
                        c.fill = ok_fill

                wb.save(filepath)
                hide_file(filepath)
            except Exception as e:
                print(f"[엑셀 단품 기록 실패]: {e}")

    def direct_finalize_excel_group(self, model_name, pallet_code, box_qr, box_time, items, header_text):
        with self.file_lock:
            try:
                filename = get_quarter_filename(model_name)
                filepath = os.path.join(BASE_DIR, filename)
                wb, ws = self.open_or_init_workbook(filepath)

                item_codes = set(it["code"].strip().upper() for it in items)

                for row in reversed(list(ws.iter_rows(min_row=2, max_row=ws.max_row))):
                    dmc_val = str(row[4].value).strip().upper() if row[4].value else ""
                    if dmc_val in item_codes:
                        row[0].value = pallet_code
                        row[1].value = box_qr
                        row[2].value = box_time
                        item_codes.remove(dmc_val)
                    if not item_codes:
                        break

                thin_border = Border(
                    left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                    top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
                )

                header_row = [pallet_code, box_qr, box_time, "HEADER", header_text, box_time, "OK", ""]
                ws.append(header_row)
                h_row_idx = ws.max_row
                for col in range(1, 9):
                    c = ws.cell(row=h_row_idx, column=col)
                    c.border = thin_border
                    c.alignment = Alignment(horizontal="center" if col in [3, 4, 6, 7, 8] else "left", vertical="center")

                wb.save(filepath)
                hide_file(filepath)
            except Exception as e:
                print(f"[그룹핑 엑셀 기록 실패]: {e}")

    def direct_handle_dmc_duplicate(self, model_name, raw_code, day_str, time_str, matched_label, dup_text, pallet_code):
        with self.file_lock:
            try:
                clean_target = raw_code.strip().upper()
                files = self.get_all_model_files(model_name)
                current_quarter_file = os.path.join(BASE_DIR, get_quarter_filename(model_name))
                if current_quarter_file not in files:
                    files.append(current_quarter_file)

                ng_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
                ng_font = Font(name="맑은 고딕", size=10, bold=True, color="C00000")
                thin_border = Border(
                    left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                    top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
                )

                for f_path in files:
                    if not os.path.exists(f_path):
                        continue
                    unhide_file(f_path)
                    wb = openpyxl.load_workbook(f_path)
                    ws = wb["스캔실적"] if "스캔실적" in wb.sheetnames else wb.active
                    modified = False

                    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                        seq_cell = row[3]
                        dmc_cell = row[4]
                        res_cell = row[6]
                        content_cell = row[7] if len(row) >= 8 else None

                        dmc_val = str(dmc_cell.value).strip() if dmc_cell.value else ""
                        lbl_val = str(row[1].value).strip() if row[1].value else ""

                        clean_val = dmc_val.replace("[중복스캔] ", "").replace("[중복 스캔] ", "").upper()
                        if clean_val == clean_target and res_cell.value != "NG":
                            res_cell.value = "NG"
                            if content_cell:
                                content_cell.value = dup_text
                            for c in row[:8]:
                                c.fill = ng_fill
                                c.font = ng_font
                            modified = True

                        elif matched_label and lbl_val == matched_label and (str(seq_cell.value) == "HEADER" or "[" in dmc_val):
                            res_cell.value = "NG"
                            for c in row[:8]:
                                c.fill = ng_fill
                                c.font = ng_font
                            modified = True

                    if modified:
                        wb.save(f_path)
                    hide_file(f_path)

                wb_cur, ws_cur = self.open_or_init_workbook(current_quarter_file)
                ts_full = f"{day_str} {time_str}"
                row_data = [pallet_code, "-", "-", "-", raw_code, ts_full, "NG", dup_text]
                ws_cur.append(row_data)
                h_idx = ws_cur.max_row
                for col in range(1, 9):
                    c = ws_cur.cell(row=h_idx, column=col)
                    c.border = thin_border
                    c.fill = ng_fill
                    c.font = ng_font
                    c.alignment = Alignment(horizontal="center" if col in [3, 4, 6, 7, 8] else "left", vertical="center")

                wb_cur.save(current_quarter_file)
                hide_file(current_quarter_file)
            except Exception as e:
                pass

    def update_stat_cards(self):
        curr_model = self.current_model.get()
        counts = self.model_counts.get(curr_model, {"total": 0, "ok": 0, "ng": 0})
        self.lbl_total_val.config(text=str(counts["total"]))
        self.lbl_ok_val.config(text=str(counts["ok"]))
        self.lbl_ng_val.config(text=str(counts["ng"]))

    def open_reset_dialog(self):
        win = tk.Toplevel(self.root)
        win.title(self.t("reset_btn"))
        win.configure(bg=BG_PANEL)
        win.transient(self.root)
        win.grab_set()

        self.center_popup(win, 360, 210)

        curr_model = self.current_model.get()
        prompt_txt = f"[{curr_model}] " + ("카운터를 초기화하려면\n관리자 비밀번호를 입력하세요." if self.current_lang.get()=="한국어" else "Enter Admin Password to Reset Counter.")
        tk.Label(win, text=prompt_txt, font=("맑은 고딕", 10, "bold"), fg=TEXT_COLOR, bg=BG_PANEL).pack(pady=(15, 8))

        pw_entry = tk.Entry(win, show="*", font=("Arial", 14), justify="center", bg=BG_INPUT, fg="#ffffff")
        pw_entry.pack(pady=5)
        pw_entry.focus_set()

        lbl_err = tk.Label(win, text="", font=("맑은 고딕", 9), fg="#ff6b6b", bg=BG_PANEL)
        lbl_err.pack()

        def do_reset(event=None):
            if pw_entry.get() == self.admin_password:
                self.model_counts[curr_model] = {"total": 0, "ok": 0, "ng": 0}
                self.save_model_counts()
                self.update_stat_cards()
                win.destroy()
                self.scan_entry.focus_set()
            else:
                lbl_err.config(text=self.t("pw_err"))
                pw_entry.delete(0, tk.END)

        pw_entry.bind("<Return>", do_reset)
        tk.Button(win, text=self.t("unlock_btn"), command=do_reset, bg="#dc3545", fg="#ffffff",
                  relief="flat", font=("맑은 고딕", 10, "bold"), padx=15, pady=3).pack(pady=10)

    def apply_recode_filter(self):
        self.tree_recode.delete(*self.tree_recode.get_children())
        model = self.current_model.get()
        target_upper = MODEL_CONFIG[model].upper()
        files = self.get_all_model_files(model)

        if not files:
            return

        start_day = self.entry_start_day.get().strip()
        end_day = self.entry_end_day.get().strip()
        start_time = self.entry_start_time.get().strip()
        end_time = self.entry_end_time.get().strip()

        try:
            matched = []
            for filepath in files:
                unhide_file(filepath)
                wb = openpyxl.load_workbook(filepath, data_only=True)
                if "스캔실적" in wb.sheetnames:
                    ws = wb["스캔실적"]
                else:
                    sheets = [s for s in wb.worksheets if s.title != "sorting"]
                    ws = sheets[0] if sheets else wb.active

                last_known_label_qr = ""

                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or len(row) < 7:
                        continue
                    
                    if len(row) >= 8:
                        pallet_val = str(row[0]).strip().upper() if row[0] and str(row[0]).strip() != "-" else ""
                        label_qr = row[1]
                        box_time = row[2]
                        seq_val = row[3]
                        dmc_code = row[4]
                        dmc_time = row[5]
                        res = row[6]
                        content = row[7] if len(row) >= 8 and row[7] else ""
                    else:
                        pallet_val = ""
                        label_qr = row[0]
                        box_time = row[1]
                        seq_val = row[2]
                        dmc_code = row[3]
                        dmc_time = row[4]
                        res = row[5]
                        content = row[6] if len(row) >= 7 and row[6] else ""

                    if label_qr and str(label_qr).strip() != "-":
                        last_known_label_qr = str(label_qr).strip()

                    ts_str = str(dmc_time if dmc_time and dmc_time != "-" else box_time)
                    if not ts_str or ts_str == "-":
                        continue

                    parts = ts_str.split()
                    r_day = parts[0] if len(parts) > 0 else ""
                    r_time = parts[1] if len(parts) > 1 else "00:00:00"

                    if start_day and r_day < start_day:
                        continue
                    if end_day and r_day > end_day:
                        continue
                    if start_time and r_time < start_time:
                        continue
                    if end_time and r_time > end_time:
                        continue

                    lbl_val = last_known_label_qr
                    dmc_val = "" if not dmc_code or dmc_code == "-" else str(dmc_code)
                    content_val = str(content).strip() if content else ""

                    clean_dmc = dmc_val.replace("[중복스캔] ", "").replace("[중복 스캔] ", "").upper()
                    if not (clean_dmc.startswith(target_upper) or lbl_val.upper().startswith(target_upper) or "PALLET" in dmc_val):
                        continue

                    matched.append((pallet_val, r_day, r_time, lbl_val, dmc_val, str(res), content_val))

                hide_file(filepath)

            for m in reversed(matched):
                tag = ""
                if m[5] == "NG":
                    tag = "ng_row"
                elif "Pallet" in str(m[4]):
                    tag = "pallet_row"
                self.tree_recode.insert("", tk.END, values=m, tags=(tag,) if tag else ())

        except Exception:
            pass

    def save_recode_to_excel(self):
        items = self.tree_recode.get_children()
        if not items:
            return

        now_ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        initial_name = f"{self.current_model.get()}_Recode_{now_ts}.xlsx"
        save_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel Files", "*.xlsx")],
            initialfile=initial_name,
            title="Save Records"
        )

        if not save_path:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Re-code"

            headers = ["Pallet Label QR", "DAY", "TIME", "Label QR", "DMC", "판정", "Content"]
            ws.append(headers)

            header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
            header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
            thin_border = Border(
                left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
            )

            for col in range(1, 8):
                cell = ws.cell(row=1, column=col)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for idx, item_id in enumerate(items, start=2):
                row_vals = self.tree_recode.item(item_id, "values")
                ws.append(list(row_vals))
                for col in range(1, 8):
                    c = ws.cell(row=idx, column=col)
                    c.border = thin_border
                    c.alignment = Alignment(horizontal="center" if col in [2, 3, 6, 7] else "left", vertical="center")

            ws.column_dimensions['A'].width = 30
            ws.column_dimensions['B'].width = 14
            ws.column_dimensions['C'].width = 14
            ws.column_dimensions['D'].width = 46
            ws.column_dimensions['E'].width = 34
            ws.column_dimensions['F'].width = 12
            ws.column_dimensions['G'].width = 16

            wb.save(save_path)
            messagebox.showinfo("Success", f"Saved successfully:\n{save_path}")

        except Exception:
            pass

    def open_lock_popup(self, title_text, msg, header_bg, header_fg):
        self.play_alarm_sound()

        dialog = tk.Toplevel(self.root)
        dialog.title(title_text)
        dialog.resizable(False, False)
        dialog.configure(bg=header_bg)

        dialog.transient(self.root)
        dialog.grab_set()
        dialog.protocol("WM_DELETE_WINDOW", lambda: None)

        self.center_popup(dialog, 520, 320)
        self.active_popup = dialog

        tk.Label(dialog, text=msg, font=("맑은 고딕", 11), bg=header_bg, fg=header_fg, justify=tk.LEFT).pack(pady=15)

        pw_entry = tk.Entry(dialog, show="*", font=("Arial", 16), justify="center", width=14,
                            bg=BG_INPUT, fg="#ffffff", insertbackground="#ffffff")
        pw_entry.pack(pady=5)
        pw_entry.focus_set()

        lbl_err = tk.Label(dialog, text="", font=("맑은 고딕", 10, "bold"), fg="#ff6b6b", bg=header_bg)
        lbl_err.pack()

        def unlock(event=None):
            if pw_entry.get() == self.admin_password:
                dialog.grab_release()
                dialog.destroy()
                self.active_popup = None
                self.set_status("READY", "#adb5bd", "#2a2e37")
                self.scan_entry.focus_set()
            else:
                lbl_err.config(text=self.t("pw_err"))
                pw_entry.delete(0, tk.END)

        pw_entry.bind("<Return>", unlock)
        tk.Button(dialog, text=self.t("unlock_btn"), command=unlock,
                  font=("맑은 고딕", 11, "bold"), bg=header_fg, fg="#ffffff",
                  relief="flat", padx=16, pady=5, cursor="hand2").pack(pady=12)

    def change_password_dialog(self):
        win = tk.Toplevel(self.root)
        win.title(self.t("pw_setting"))
        win.configure(bg=BG_PANEL)
        win.transient(self.root)
        win.grab_set()

        self.center_popup(win, 340, 220)

        prompt_curr = "현재 비밀번호" if self.current_lang.get()=="한국어" else "Current Password"
        prompt_new = "새 6자리 숫자 비밀번호" if self.current_lang.get()=="한국어" else "New 6-digit Password"

        tk.Label(win, text=prompt_curr, font=("맑은 고딕", 9), fg=TEXT_COLOR, bg=BG_PANEL).pack(pady=(15, 2))
        curr_entry = tk.Entry(win, show="*", font=("Arial", 11), justify="center", bg=BG_INPUT, fg="#ffffff")
        curr_entry.pack()

        tk.Label(win, text=prompt_new, font=("맑은 고딕", 9), fg=TEXT_COLOR, bg=BG_PANEL).pack(pady=(10, 2))
        new_entry = tk.Entry(win, show="*", font=("Arial", 11), justify="center", bg=BG_INPUT, fg="#ffffff")
        new_entry.pack()

        def apply_pw():
            if curr_entry.get() != self.admin_password:
                messagebox.showerror("Error", self.t("pw_err"), parent=win)
                return
            new_val = new_entry.get()
            if len(new_val) != 6 or not new_val.isdigit():
                messagebox.showerror("Error", "Password must be 6 digits.", parent=win)
                return
            self.admin_password = new_val
            messagebox.showinfo("Success", "Password changed successfully.", parent=win)
            win.destroy()

        tk.Button(win, text=self.t("unlock_btn"), command=apply_pw, bg="#2b5278", fg="#ffffff",
                  relief="flat", font=("맑은 고딕", 10, "bold"), padx=15, pady=4).pack(pady=15)


if __name__ == "__main__":
    root = tk.Tk()
    app = QRScanStationApp(root)
    root.mainloop()
